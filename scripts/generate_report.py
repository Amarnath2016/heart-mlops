"""Build reports/MLOps_Assignment_Report.docx from the real pipeline outputs.

Run after: download_data -> src.eda -> src.train, and after dropping screenshots into
screenshots/. Metrics, tables and figures are read from the artifacts, so the report
always matches what was actually run.

    python scripts/generate_report.py --repo-url https://github.com/<you>/<repo> --name "Your Name" --id 2024XX0000
"""
import argparse
import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src import config  # noqa: E402
from src.data import load_clean, load_raw  # noqa: E402

SHOTS = [
    ("mlflow_ui.png", "MLflow UI - experiment runs"),
    ("mlflow_compare.png", "MLflow - run comparison"),
    ("github_actions.png", "GitHub Actions - successful workflow run"),
    ("github_actions_artifacts.png", "GitHub Actions - uploaded artifacts"),
    ("docker_build.png", "Docker image build"),
    ("docker_predict.png", "Container running and answering /predict"),
    ("kubectl_get_all.png", "kubectl get pods,svc - deployment status"),
    ("k8s_predict.png", "/predict via the Kubernetes LoadBalancer"),
    ("api_logs.png", "Structured request logs"),
    ("prometheus_targets.png", "Prometheus target UP"),
    ("grafana_dashboard.png", "Grafana dashboard"),
]


def architecture_diagram(path):
    fig, ax = plt.subplots(figsize=(11, 4.2))
    ax.axis("off")
    boxes = [
        (0.2, 2.6, "UCI Heart\nDisease data", "#DCE6F2"),
        (2.2, 2.6, "download_data.py\nclean + EDA", "#DCE6F2"),
        (4.4, 2.6, "train.py\nCV + GridSearch\n(LR, RF)", "#E2F0D9"),
        (6.6, 2.6, "MLflow tracking\nparams / metrics /\nplots / model", "#FFF2CC"),
        (8.8, 2.6, "models/model.joblib\n(preprocess + clf)", "#E2F0D9"),
        (2.2, 0.4, "GitHub Actions\nlint > test > train\n> docker build", "#F8CBAD"),
        (4.9, 0.4, "Docker image\nFlask + gunicorn", "#F8CBAD"),
        (7.2, 0.4, "Kubernetes\nDeployment + Service\n(LoadBalancer)", "#F8CBAD"),
        (9.2, 0.4, "Prometheus +\nGrafana, JSON logs", "#E4DFEC"),
    ]
    for x, y, t, c in boxes:
        ax.add_patch(FancyBboxPatch((x, y), 1.7, 1.2, boxstyle="round,pad=0.05", fc=c, ec="#444"))
        ax.text(x + 0.85, y + 0.6, t, ha="center", va="center", fontsize=8)
    arrows = [((1.9, 3.2), (2.2, 3.2)), ((3.9, 3.2), (4.4, 3.2)), ((6.1, 3.2), (6.6, 3.2)),
              ((8.3, 3.2), (8.8, 3.2)), ((3.05, 2.6), (3.05, 1.6)), ((3.9, 1.0), (4.9, 1.0)),
              ((6.6, 1.0), (7.2, 1.0)), ((8.9, 1.0), (9.2, 1.0)), ((9.65, 2.6), (6.4, 1.6))]
    for a, b in arrows:
        ax.annotate("", xy=b, xytext=a, arrowprops=dict(arrowstyle="->", color="#333"))
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 4.2)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def add_table(doc, rows, header):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)


def fig(doc, path, caption, width=6.0):
    path = Path(path)
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.runs[0].italic = True
    else:
        doc.add_paragraph(f"[MISSING: {path.name} - {caption}]")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo-url", default="<add GitHub repository URL>")
    ap.add_argument("--name", default="")
    ap.add_argument("--id", default="")
    ap.add_argument("--out", default=str(config.ROOT / "reports" / "MLOps_Assignment_Report.docx"))
    a = ap.parse_args()

    meta = json.loads(config.METADATA_PATH.read_text())
    clean, raw = load_clean(), load_raw()
    miss = raw.isna().sum()
    miss = miss[miss > 0]
    corr = clean.corr(numeric_only=True)[config.TARGET].drop(config.TARGET)
    top = corr.reindex(corr.abs().sort_values(ascending=False).index).head(5)
    vc = clean[config.TARGET].value_counts()
    diagram = config.ROOT / "reports" / "figures" / "architecture.png"
    architecture_diagram(diagram)

    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading("MLOps Assignment 1 - Heart Disease Risk Prediction", 0)
    d.add_paragraph(f"{a.name}  {a.id}\nCourse: Machine Learning Operations (AIMLCZG523)\n"
                    f"Code repository: {a.repo_url}")

    d.add_heading("1. Setup and installation", 1)
    d.add_paragraph("Python 3.11 is used. From a clean checkout:")
    for line in ["python -m venv .venv && source .venv/bin/activate",
                 "pip install -r requirements.txt",
                 "python scripts/download_data.py        # downloads UCI data, writes data/processed/heart_clean.csv",
                 "python -m src.eda                      # figures in reports/figures",
                 "python -m src.train                    # MLflow runs in ./mlruns, model in models/",
                 "mlflow ui --backend-store-uri ./mlruns # browse experiments",
                 "pytest -v",
                 "docker build -t heart-api . && docker run -p 8000:8000 heart-api",
                 "kubectl apply -f k8s/"]:
        d.add_paragraph(line, style="List Bullet")

    d.add_heading("2. Data and exploratory analysis", 1)
    d.add_paragraph(
        f"The Cleveland subset of the UCI Heart Disease dataset has {raw.shape[0]} rows and 13 predictors plus a "
        f"0-4 diagnosis column. The diagnosis was binarised (0 = no disease, 1-4 = disease). After removing "
        f"duplicates, {len(clean)} rows remain: {int(vc.get(0, 0))} without and {int(vc.get(1, 0))} with disease "
        f"({vc.get(1, 0) / len(clean):.1%} positive), so the classes are reasonably balanced and accuracy is "
        f"meaningful alongside recall and ROC-AUC.")
    if len(miss):
        d.add_paragraph("Missing values ('?' in the raw file): " +
                        ", ".join(f"{k} ({int(v)})" for k, v in miss.items()) +
                        ". They are imputed inside the modelling pipeline (median for numeric, most-frequent for "
                        "categorical) so that statistics are learned on training folds only.")
    d.add_paragraph("Features most correlated with the target (Pearson, |r|): " +
                    ", ".join(f"{k} ({v:+.2f})" for k, v in top.items()) + ".")
    F = config.FIGURES_DIR
    fig(d, F / "class_balance.png", "Figure 1 - Class balance", 3.2)
    fig(d, F / "numeric_histograms.png", "Figure 2 - Numeric feature distributions by target")
    fig(d, F / "correlation_heatmap.png", "Figure 3 - Correlation heatmap", 5.2)
    fig(d, F / "categorical_vs_target.png", "Figure 4 - Categorical features vs target")

    d.add_heading("3. Feature engineering and model development", 1)
    d.add_paragraph(
        "A single scikit-learn ColumnTransformer is used for both training and serving: numeric features "
        "(age, trestbps, chol, thalach, oldpeak) are median-imputed and standardised; categorical codes (sex, cp, "
        "fbs, restecg, exang, slope, ca, thal) are mode-imputed and one-hot encoded with handle_unknown='ignore'. "
        "Because the transformer is part of the model pipeline, the API only needs the raw 13 fields.")
    d.add_paragraph(
        "Data was split 80/20 (stratified, seed 42). Logistic Regression and Random Forest were tuned with "
        "GridSearchCV (5-fold stratified CV, ROC-AUC as the objective) on the training split only. The final model "
        "is chosen by mean cross-validated ROC-AUC; the hold-out test set is used only for the final report.")
    rows = [(n, f"{v['cv_roc_auc_mean']:.3f}", f"{v['test_accuracy']:.3f}", f"{v['test_precision']:.3f}",
             f"{v['test_recall']:.3f}", f"{v['test_roc_auc']:.3f}") for n, v in meta["comparison"].items()]
    add_table(d, rows, ["Model", "CV ROC-AUC", "Test acc.", "Test precision", "Test recall", "Test ROC-AUC"])
    d.add_paragraph(f"\nSelected model: {meta['model_name']} with parameters {meta['best_params']}.")
    cvm = meta["cv_metrics"]
    d.add_paragraph("Cross-validated metrics of the selected model (mean +/- std): " + "; ".join(
        f"{m} {cvm[f'cv_{m}_mean']:.3f} +/- {cvm[f'cv_{m}_std']:.3f}"
        for m in ["accuracy", "precision", "recall", "f1", "roc_auc"]) + ".")
    d.add_paragraph("Note: with ~240 training rows, test-set metrics have wide uncertainty; differences of a "
                    "few points between models are not conclusive.")

    d.add_heading("4. Experiment tracking", 1)
    d.add_paragraph(
        "Each model family is one MLflow run in experiment 'heart-disease-risk'. Logged: best hyper-parameters, grid "
        "size, split sizes, CV mean/std and hold-out metrics, ROC curve and confusion-matrix plots, and the fitted "
        f"pipeline as an MLflow model. Selected run id: {meta['mlflow_run_id']}.")

    d.add_heading("5. Packaging and reproducibility", 1)
    d.add_paragraph("The full pipeline (preprocessing + classifier) is saved with joblib to models/model.joblib "
                    "with models/model_metadata.json (features, params, metrics, run id). Dependencies are pinned in "
                    "requirements.txt; the serving image uses requirements-serve.txt with identical versions. Random "
                    "seeds are fixed.")

    d.add_heading("6. CI/CD and automated testing", 1)
    d.add_paragraph(
        "Unit tests (tests/) cover data parsing and cleaning, preprocessing output, stratified splitting, unseen "
        "categories, metric ranges, input validation and the API (health, predict, bad input, metrics). The GitHub "
        "Actions workflow runs flake8, pytest, data download + EDA + training, uploads model/MLflow/reports as "
        "artifacts, then builds the Docker image and smoke-tests /health and /predict. Any failing step fails "
        "the run with its logs.")

    d.add_heading("7. Containerisation", 1)
    d.add_paragraph("The Flask API runs under gunicorn in a python:3.11-slim image as a non-root user. "
                    "POST /predict accepts the 13 fields as JSON and returns prediction, label, confidence and "
                    "probability of heart disease. Other endpoints: /health, /model-info, /metrics.")
    d.add_paragraph('{"age":63,"sex":1,"cp":4,"trestbps":145,"chol":233,"fbs":1,"restecg":2,"thalach":150,'
                    '"exang":0,"oldpeak":2.3,"slope":3,"ca":0,"thal":6}', style="Intense Quote")

    d.add_heading("8. Production deployment", 1)
    d.add_paragraph("k8s/deployment.yaml (2 replicas, readiness/liveness probes, resource limits) and "
                    "k8s/service.yaml (type LoadBalancer) are applied to a local Kubernetes cluster; an optional "
                    "Ingress is provided. Access instructions are in the README.")

    d.add_heading("9. Monitoring and logging", 1)
    d.add_paragraph("Every request is logged as one JSON line (timestamp, request id, method, path, status, latency, "
                    "prediction). /metrics exposes request counts by endpoint/status, predictions per class and a "
                    "latency histogram; Prometheus scrapes it and a provisioned Grafana dashboard shows request rate, "
                    "p95 latency, prediction mix and error rate (monitoring/docker-compose.yml).")

    d.add_heading("10. Architecture", 1)
    fig(d, diagram, "Figure - End-to-end architecture")

    d.add_heading("11. Screenshots", 1)
    for fn, cap in SHOTS:
        fig(d, config.ROOT / "screenshots" / fn, cap)

    d.add_heading("12. Repository", 1)
    d.add_paragraph(a.repo_url)
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    d.save(a.out)
    print("Saved", a.out)


if __name__ == "__main__":
    main()
